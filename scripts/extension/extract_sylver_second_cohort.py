import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
D = ROOT / 'method_screen/DIPK'
S = ROOT / 'method_screen/SYLVER'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def read_csv(p):
    with p.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def main():
    header = D / 'inference_feasibility/GSE25065_matrix_header.txt'
    workbook = S / 'source_figure4.xlsx'
    reference = D / 'literature/RPS_2017_supplementary_tables.xlsx'
    assert sha(workbook) == '63484d3a3e9f0b7f31e9d8183775b4cd7100520804c7711b1228e1b2bc89868f'
    assert sha(reference) == '262f38bee76d73369d6927e328c742aaaf7f5e5a54e1e9a4de44ab2b9532c375'
    fields = {}
    for row in csv.reader(header.read_text().splitlines(), delimiter='\t'):
        if not row or not row[0].startswith('!Sample_'):
            continue
        if row[0] == '!Sample_characteristics_ch1':
            key = row[1].partition(': ')[0]
            fields[key] = [v.partition(': ')[2] for v in row[1:]]
        else:
            fields[row[0][8:]] = row[1:]
    meta = [dict(zip(fields, v)) for v in zip(*fields.values())]
    assert len(meta) == 198
    assert all(len(v) == 198 for v in fields.values())
    by_id = {r['sample id']: r for r in meta}
    assert len(by_id) == 198
    projected = {r['GSM']: r for r in read_csv(D / 'inference_feasibility/GSE25065_candidate_metadata.csv')}
    prior = read_csv(ROOT / 'GSE25055_metadata.csv')
    third = read_csv(D / 'third_cohort/GSE20194_verified_array_crosswalk.csv')
    rps = openpyxl.load_workbook(reference, read_only=True, data_only=True)
    flagged = {str(r[0]) for r in rps['TableS3'].values if str(r[0]).startswith('GSM')}
    w = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    rows = []
    for values in w['g'].iter_rows(min_row=3, max_row=183, min_col=6, max_col=9, values_only=True):
        source_id, score, outcome, receptor = values
        assert source_id.startswith('H.')
        r = by_id[source_id[2:]]
        assert r['title'] == source_id[2:]
        label = outcome.split(' (', 1)[0]
        assert label == r['pathologic_response_pcr_rd']
        rows.append(dict(GSM=r['geo_accession'], source_sample_id=source_id, sample_id=r['sample id'],
                         source=r['source'], patient_key=projected[r['geo_accession']]['patient_key'],
                         SLMHRD=score, published_outcome=outcome, pathology=int(label == 'pCR'),
                         proxy=int(r['dlda30_prediction'] == 'pCR'), published_receptor=receptor,
                         rps_flag=r['geo_accession'] in flagged))
    assert len(rows) == len({r['GSM'] for r in rows}) == len({r['patient_key'] for r in rows}) == 181
    assert all(by_id[r['sample_id']]['dlda30_prediction'] in ['pCR', 'RD'] for r in rows)
    norm = lambda x: re.sub(r'^M(?=\d+$)', '', x)
    prior_keys = {norm(r['sample id']) for r in prior if r['source'] == 'MDACC'}
    third_keys = {norm(r['patient_key'].split(':', 1)[1]) for r in third if r['patient_key'].startswith('MDACC:')}
    mdacc = [r for r in meta if r['source'] == 'MDACC']
    overlaps = {name: [r['geo_accession'] for r in mdacc if norm(r['sample id']) in keys]
                for name, keys in [('GSE25055_MDACC_alias', prior_keys), ('GSE20194_MDACC_alias', third_keys)]}
    raw = lambda s: re.sub(r'^GSM\d+_', '', s.rsplit('/', 1)[-1]).casefold()
    for name, comparison in [('GSE25055', prior), ('GSE20194', third)]:
        prev_raw = {raw(r.get('raw_basename') or r['supplementary_file']) for r in comparison}
        overlaps[name + '_raw_basename'] = [r['GSM'] for r in projected.values() if raw(r['raw_basename']) in prev_raw]
    assert not any(overlaps.values()), overlaps
    out = S / 'hatzis2_released_scores.csv'
    assert not out.exists(), 'Do not overwrite frozen extracted scores'
    with out.open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys()); writer.writeheader(); writer.writerows(rows)
    missing = [r for r in meta if r['pathologic_response_pcr_rd'] in ['RD', 'pCR'] and r['geo_accession'] not in {x['GSM'] for x in rows}]
    assert [r['geo_accession'] for r in missing] == ['GSM615791']
    print(json.dumps({'rows':len(rows), 'rps_flagged':sum(r['rps_flag'] for r in rows), 'overlaps':overlaps, 'missing':{k: missing[0][k] for k in ['geo_accession', 'sample id', 'pathologic_response_pcr_rd', 'dlda30_prediction']}}))


if __name__ == '__main__':
    main()
